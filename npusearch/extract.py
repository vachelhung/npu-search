"""Plain-text extraction for the file types in config.json."""
import json
import re
from pathlib import Path


def _pdf(p: Path) -> str:
    import pymupdf
    with pymupdf.open(p) as doc:
        return "\n".join(page.get_text() for page in doc)


def _docx(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xlsx(p: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _pptx(p: Path) -> str:
    import pptx
    parts = []
    for i, slide in enumerate(pptx.Presentation(str(p)).slides, 1):
        parts.append(f"# Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _ipynb(p: Path) -> str:
    # Cell sources only; outputs are mostly images and logs.
    nb = json.loads(p.read_text(encoding="utf-8", errors="replace"))
    return "\n\n".join("".join(c.get("source", [])) for c in nb.get("cells", []))


def _rtf(p: Path) -> str:
    raw = p.read_text(encoding="latin-1", errors="replace")
    return re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", raw)


def _text(p: Path) -> str:
    data = p.read_bytes()
    for enc in ("utf-8", "cp950", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


HANDLERS = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".pptx": _pptx,
            ".ipynb": _ipynb, ".rtf": _rtf}


def extract(p: Path, max_chars: int) -> str:
    text = HANDLERS.get(p.suffix.lower(), _text)(p)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    return text[:max_chars]
